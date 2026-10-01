from datetime import date
from io import BytesIO

from django.db import connection
from django.http import HttpResponse
from django.shortcuts import render, redirect

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from bantaybarangay.security import role_required

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)


# =========================================================
# DATABASE HELPERS
# =========================================================

def fetch_one(query, params=None):

    with connection.cursor() as cursor:

        cursor.execute(
            query,
            params or []
        )

        row = cursor.fetchone()

    if not row:
        return 0

    return row[0] or 0


def fetch_all_dict(query, params=None):

    with connection.cursor() as cursor:

        cursor.execute(
            query,
            params or []
        )

        columns = [
            column[0]
            for column in cursor.description
        ]

        return [
            dict(zip(columns, row))
            for row in cursor.fetchall()
        ]





# =========================================================
# REPORT DATA
# =========================================================

def get_report_data():

    today = date.today()

    current_year = today.year
    current_month = today.month


    # =====================================================
    # SUMMARY
    # =====================================================

    monthly_complaints = fetch_one(
        """
        SELECT COUNT(*)
        FROM complaints
        WHERE YEAR(submitted_at) = %s
          AND MONTH(submitted_at) = %s
        """,
        [
            current_year,
            current_month,
        ]
    )


    incident_reports = fetch_one(
        """
        SELECT COUNT(*)
        FROM complaints
        WHERE report_type = %s
          AND YEAR(submitted_at) = %s
          AND MONTH(submitted_at) = %s
        """,
        [
            "Community Issue",
            current_year,
            current_month,
        ]
    )


    total_complaints = fetch_one(
        """
        SELECT COUNT(*)
        FROM complaints
        """
    )


    resolved_complaints = fetch_one(
        """
        SELECT COUNT(*)
        FROM complaints
        WHERE status IN (
            'Settled',
            'Resolved',
            'Closed'
        )
        """
    )


    if total_complaints > 0:

        resolution_rate = round(
            (
                resolved_complaints /
                total_complaints
            ) * 100,
            1
        )

    else:
        resolution_rate = 0


    released_documents = fetch_one(
        """
        SELECT COUNT(*)
        FROM document_requests
        WHERE status = %s
        """,
        [
            "Released"
        ]
    )


    # =====================================================
    # MONTHLY COMPLAINT TREND
    # =====================================================

    monthly_rows = fetch_all_dict(
        """
        SELECT
            MONTH(submitted_at) AS month_number,
            COUNT(*) AS total
        FROM complaints
        WHERE YEAR(submitted_at) = %s
        GROUP BY MONTH(submitted_at)
        ORDER BY MONTH(submitted_at)
        """,
        [
            current_year
        ]
    )


    month_counts = {
        row["month_number"]:
            row["total"]
        for row in monthly_rows
    }


    month_names = [
        "JAN",
        "FEB",
        "MAR",
        "APR",
        "MAY",
        "JUN",
        "JUL",
        "AUG",
        "SEP",
        "OCT",
        "NOV",
        "DEC",
    ]


    complaint_chart_labels = []
    complaint_chart_values = []


    for month_number in range(
        1,
        13
    ):

        complaint_chart_labels.append(
            month_names[
                month_number - 1
            ]
        )

        complaint_chart_values.append(
            month_counts.get(
                month_number,
                0
            )
        )


    # =====================================================
    # COMPLAINT CATEGORIES
    # =====================================================

    category_rows = fetch_all_dict(
        """
        SELECT
            complaint_type,
            COUNT(*) AS total
        FROM complaints
        GROUP BY complaint_type
        ORDER BY total DESC
        """
    )


    category_total = sum(
        row["total"]
        for row in category_rows
    )


    complaint_categories = []


    for row in category_rows:

        if category_total > 0:

            percentage = round(
                (
                    row["total"] /
                    category_total
                ) * 100,
                1
            )

        else:

            percentage = 0


        complaint_categories.append(
            {
                "name":
                    row["complaint_type"],

                "total":
                    row["total"],

                "percentage":
                    percentage,
            }
        )


    category_chart_labels = [
        item["name"]
        for item in complaint_categories
    ]


    category_chart_values = [
        item["total"]
        for item in complaint_categories
    ]


    # =====================================================
    # RESOLUTION TREND
    # =====================================================

    resolution_rows = fetch_all_dict(
        """
        SELECT
            MONTH(submitted_at) AS month_number,

            COUNT(*) AS filed,

            SUM(
                CASE
                    WHEN status IN (
                        'Settled',
                        'Resolved',
                        'Closed'
                    )
                    THEN 1
                    ELSE 0
                END
            ) AS resolved

        FROM complaints

        WHERE YEAR(submitted_at) = %s

        GROUP BY MONTH(submitted_at)

        ORDER BY MONTH(submitted_at)
        """,
        [
            current_year
        ]
    )


    resolution_lookup = {
        row["month_number"]: row
        for row in resolution_rows
    }


    resolution_labels = []
    resolution_filed = []
    resolution_resolved = []


    for month_number in range(
        1,
        13
    ):

        resolution_labels.append(
            month_names[
                month_number - 1
            ]
        )


        row = resolution_lookup.get(
            month_number
        )


        if row:

            resolution_filed.append(
                row["filed"] or 0
            )

            resolution_resolved.append(
                row["resolved"] or 0
            )

        else:

            resolution_filed.append(0)
            resolution_resolved.append(0)


    # =====================================================
    # INCIDENT CATEGORIES
    # =====================================================

    incident_rows = fetch_all_dict(
        """
        SELECT
            complaint_type,
            COUNT(*) AS total

        FROM complaints

        WHERE report_type = %s

        GROUP BY complaint_type

        ORDER BY total DESC

        LIMIT 4
        """,
        [
            "Community Issue"
        ]
    )


    maximum_incident = max(
        [
            row["total"]
            for row in incident_rows
        ] + [1]
    )


    incident_categories = []


    for row in incident_rows:

        progress = round(
            (
                row["total"] /
                maximum_incident
            ) * 100
        )


        incident_categories.append(
            {
                "name":
                    row["complaint_type"],

                "total":
                    row["total"],

                "progress":
                    progress,
            }
        )


    # =====================================================
    # EVIDENCE
    # =====================================================

    total_evidence = fetch_one(
        """
        SELECT COUNT(*)
        FROM evidence
        """
    )


    registered_evidence = fetch_one(
        """
        SELECT COUNT(*)
        FROM evidence
        WHERE blockchain_status = %s
        """,
        [
            "Registered"
        ]
    )


    pending_evidence = fetch_one(
        """
        SELECT COUNT(*)
        FROM evidence
        WHERE blockchain_status = %s
        """,
        [
            "Pending"
        ]
    )


    failed_evidence = fetch_one(
        """
        SELECT COUNT(*)
        FROM evidence
        WHERE blockchain_status = %s
        """,
        [
            "Failed"
        ]
    )


    if total_evidence > 0:

        blockchain_rate = round(
            (
                registered_evidence /
                total_evidence
            ) * 100,
            1
        )

    else:

        blockchain_rate = 0


    # =====================================================
    # OFFICIAL ACTIVITY
    # =====================================================

    official_activity = fetch_all_dict(
        """
        SELECT
            u.user_id,
            u.username,
            u.role,
            u.is_active,

            COUNT(a.audit_id)
                AS activity_count,

            MAX(a.created_at)
                AS last_activity

        FROM users u

        LEFT JOIN audit_logs a
            ON a.user_id = u.user_id

        WHERE u.role IN (
            'official',
            'admin'
        )

        GROUP BY
            u.user_id,
            u.username,
            u.role,
            u.is_active

        ORDER BY
            activity_count DESC,
            u.username ASC

        LIMIT 5
        """
    )


    return {

        "monthly_complaints":
            monthly_complaints,

        "incident_reports":
            incident_reports,

        "resolution_rate":
            resolution_rate,

        # Keep existing template variable working
        "verified_documents":
            released_documents,

        "released_documents":
            released_documents,

        "resolved_complaints":
            resolved_complaints,

        "total_complaints":
            total_complaints,


        "complaint_chart_labels":
            complaint_chart_labels,

        "complaint_chart_values":
            complaint_chart_values,


        "complaint_categories":
            complaint_categories,

        "category_chart_labels":
            category_chart_labels,

        "category_chart_values":
            category_chart_values,

        "category_total":
            category_total,


        "resolution_labels":
            resolution_labels,

        "resolution_filed":
            resolution_filed,

        "resolution_resolved":
            resolution_resolved,


        "incident_categories":
            incident_categories,


        "total_evidence":
            total_evidence,

        "registered_evidence":
            registered_evidence,

        "pending_evidence":
            pending_evidence,

        "failed_evidence":
            failed_evidence,

        "blockchain_rate":
            blockchain_rate,


        "official_activity":
            official_activity,


        "current_year":
            current_year,

        "generated_date":
            today,
    }


# =========================================================
# REPORT PAGE
# =========================================================

@role_required("admin", "official")
def reports(request):

    context = get_report_data()

    return render(
        request,
        "reportsmodule/reports.html",
        context
    )


# =========================================================
# EXCEL EXPORT
# =========================================================
@role_required("admin", "official")
def export_reports_excel(request):
    data = get_report_data()

    workbook = Workbook()

    


    data = get_report_data()


    workbook = Workbook()


    # =====================================================
    # SUMMARY SHEET
    # =====================================================

    summary = workbook.active

    summary.title = "Summary"


    summary.merge_cells(
        "A1:D1"
    )

    summary["A1"] = (
        "BantayBarangay "
        "Reports & Analytics"
    )

    summary["A1"].font = Font(
        bold=True,
        size=18,
        color="FFFFFF"
    )

    summary["A1"].fill = PatternFill(
        "solid",
        fgColor="064E3B"
    )

    summary["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    summary.row_dimensions[1].height = 30


    summary.merge_cells(
        "A2:D2"
    )

    summary["A2"] = (
        "Generated: "
        f"{data['generated_date']:%B %d, %Y}"
    )

    summary["A2"].alignment = Alignment(
        horizontal="center"
    )


    summary.append(
        []
    )


    summary.append(
        [
            "Metric",
            "Value",
        ]
    )


    summary_rows = [

        [
            "Monthly Complaints",
            data["monthly_complaints"],
        ],

        [
            "Monthly Incident Reports",
            data["incident_reports"],
        ],

        [
            "Total Complaints",
            data["total_complaints"],
        ],

        [
            "Resolved Cases",
            data["resolved_complaints"],
        ],

        [
            "Resolution Rate",
            f"{data['resolution_rate']}%",
        ],

        [
            "Released Documents",
            data["released_documents"],
        ],

        [
            "Total Evidence",
            data["total_evidence"],
        ],

        [
            "Registered Evidence",
            data["registered_evidence"],
        ],

        [
            "Pending Evidence",
            data["pending_evidence"],
        ],

        [
            "Failed Evidence",
            data["failed_evidence"],
        ],

        [
            "Blockchain Registration Rate",
            f"{data['blockchain_rate']}%",
        ],
    ]


    for row in summary_rows:

        summary.append(
            row
        )


    # =====================================================
    # MONTHLY COMPLAINTS SHEET
    # =====================================================

    monthly_sheet = (
        workbook.create_sheet(
            "Monthly Complaints"
        )
    )


    monthly_sheet.append(
        [
            "Month",
            "Complaints",
        ]
    )


    for label, value in zip(
        data["complaint_chart_labels"],
        data["complaint_chart_values"]
    ):

        monthly_sheet.append(
            [
                label,
                value,
            ]
        )


    # =====================================================
    # CATEGORIES SHEET
    # =====================================================

    category_sheet = (
        workbook.create_sheet(
            "Complaint Categories"
        )
    )


    category_sheet.append(
        [
            "Category",
            "Cases",
            "Percentage",
        ]
    )


    for item in data[
        "complaint_categories"
    ]:

        category_sheet.append(
            [
                item["name"],
                item["total"],
                item["percentage"] / 100,
            ]
        )


    # =====================================================
    # RESOLUTION SHEET
    # =====================================================

    resolution_sheet = (
        workbook.create_sheet(
            "Resolution Trend"
        )
    )


    resolution_sheet.append(
        [
            "Month",
            "Filed",
            "Resolved",
        ]
    )


    for (
        label,
        filed,
        resolved
    ) in zip(
        data["resolution_labels"],
        data["resolution_filed"],
        data["resolution_resolved"]
    ):

        resolution_sheet.append(
            [
                label,
                filed,
                resolved,
            ]
        )


    # =====================================================
    # INCIDENT SHEET
    # =====================================================

    incident_sheet = (
        workbook.create_sheet(
            "Community Issues"
        )
    )


    incident_sheet.append(
        [
            "Category",
            "Cases",
        ]
    )


    for item in data[
        "incident_categories"
    ]:

        incident_sheet.append(
            [
                item["name"],
                item["total"],
            ]
        )


    # =====================================================
    # OFFICIAL ACTIVITY SHEET
    # =====================================================

    official_sheet = (
        workbook.create_sheet(
            "Official Activity"
        )
    )


    official_sheet.append(
        [
            "Username",
            "Role",
            "Status",
            "Activities",
            "Last Activity",
        ]
    )


    for officer in data[
        "official_activity"
    ]:

        last_activity = (
            officer["last_activity"]
            if officer["last_activity"]
            else ""
        )

        official_sheet.append(
            [
                officer["username"],
                str(
                    officer["role"]
                ).title(),
                (
                    "Active"
                    if officer["is_active"]
                    else "Inactive"
                ),
                officer["activity_count"],
                last_activity,
            ]
        )


    # =====================================================
    # FORMAT WORKBOOK
    # =====================================================

    header_fill = PatternFill(
        "solid",
        fgColor="0F5132"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D1D5DB"
        )
    )


    for sheet in workbook.worksheets:

        if sheet.title == "Summary":

            header_row = 4

        else:

            header_row = 1


        for cell in sheet[
            header_row
        ]:

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


        for row in sheet.iter_rows():

            for cell in row:

                cell.border = thin_border

                cell.alignment = Alignment(
                    vertical="top"
                )


        for column_cells in (
            sheet.columns
        ):

            max_length = 0

            column_letter = (
                get_column_letter(
                    column_cells[0].column
                )
            )


            for cell in column_cells:

                value = cell.value

                if value is None:
                    continue

                max_length = max(
                    max_length,
                    len(str(value))
                )


            sheet.column_dimensions[
                column_letter
            ].width = min(
                max(
                    max_length + 3,
                    12
                ),
                40
            )


        sheet.freeze_panes = (
            f"A{header_row + 1}"
        )


    # Percentage format

    for row_number in range(
        2,
        category_sheet.max_row + 1
    ):

        category_sheet.cell(
            row=row_number,
            column=3
        ).number_format = "0.0%"


    # =====================================================
    # RESPONSE
    # =====================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)


    filename = (
        "bantaybarangay_reports_"
        f"{data['generated_date']:%Y-%m-%d}"
        ".xlsx"
    )


    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/"
            "vnd.openxmlformats-"
            "officedocument."
            "spreadsheetml.sheet"
        )
    )


    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )


    return response


# =========================================================
# PDF EXPORT
# =========================================================
@role_required("admin", "official")
def export_reports_pdf(request):
    data = get_report_data()

    buffer = BytesIO()
    if not request.session.get(
        "is_logged_in"
    ):

        return redirect(
            "login"
        )


    if not has_report_access(
        request
    ):

        return redirect(
            "home"
        )


    data = get_report_data()


    buffer = BytesIO()


    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=(
            "BantayBarangay "
            "Reports & Analytics"
        ),
        author="BantayBarangay"
    )


    styles = getSampleStyleSheet()


    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor(
            "#064E3B"
        ),
        spaceAfter=5 * mm,
    )


    section_style = ParagraphStyle(
        "SectionTitle",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        textColor=colors.HexColor(
            "#064E3B"
        ),
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
    )


    normal_style = styles[
        "BodyText"
    ]


    story = []


    # =====================================================
    # TITLE
    # =====================================================

    story.append(
        Paragraph(
            (
                "BantayBarangay "
                "Reports &amp; Analytics"
            ),
            title_style
        )
    )


    story.append(
        Paragraph(
            (
                "Generated on "
                f"{data['generated_date']:%B %d, %Y}"
            ),
            normal_style
        )
    )


    story.append(
        Spacer(
            1,
            5 * mm
        )
    )


    # =====================================================
    # SUMMARY
    # =====================================================

    story.append(
        Paragraph(
            "Summary",
            section_style
        )
    )


    summary_data = [

        [
            "Metric",
            "Value",
        ],

        [
            "Monthly Complaints",
            data["monthly_complaints"],
        ],

        [
            "Monthly Incident Reports",
            data["incident_reports"],
        ],

        [
            "Total Complaints",
            data["total_complaints"],
        ],

        [
            "Resolved Cases",
            data["resolved_complaints"],
        ],

        [
            "Resolution Rate",
            f"{data['resolution_rate']}%",
        ],

        [
            "Released Documents",
            data["released_documents"],
        ],

        [
            "Blockchain Registration Rate",
            f"{data['blockchain_rate']}%",
        ],
    ]


    summary_table = Table(
        summary_data,
        colWidths=[
            90 * mm,
            45 * mm,
        ],
        repeatRows=1,
        hAlign="LEFT"
    )


    apply_pdf_table_style(
        summary_table
    )


    story.append(
        summary_table
    )


    # =====================================================
    # MONTHLY COMPLAINTS
    # =====================================================

    story.append(
        Paragraph(
            (
                f"Monthly Complaints - "
                f"{data['current_year']}"
            ),
            section_style
        )
    )


    monthly_table_data = [
        [
            "Month",
            "Complaints",
        ]
    ]


    for label, value in zip(
        data["complaint_chart_labels"],
        data["complaint_chart_values"]
    ):

        monthly_table_data.append(
            [
                label,
                value,
            ]
        )


    monthly_table = Table(
        monthly_table_data,
        colWidths=[
            40 * mm,
            40 * mm,
        ],
        repeatRows=1,
        hAlign="LEFT"
    )


    apply_pdf_table_style(
        monthly_table
    )


    story.append(
        monthly_table
    )


    # =====================================================
    # CATEGORIES
    # =====================================================

    story.append(
        Paragraph(
            "Complaint Categories",
            section_style
        )
    )


    category_table_data = [
        [
            "Category",
            "Cases",
            "Percentage",
        ]
    ]


    for item in data[
        "complaint_categories"
    ]:

        category_table_data.append(
            [
                item["name"],
                item["total"],
                f"{item['percentage']}%",
            ]
        )


    category_table = Table(
        category_table_data,
        colWidths=[
            100 * mm,
            35 * mm,
            40 * mm,
        ],
        repeatRows=1,
        hAlign="LEFT"
    )


    apply_pdf_table_style(
        category_table
    )


    story.append(
        category_table
    )


    # =====================================================
    # PAGE TWO
    # =====================================================

    story.append(
        PageBreak()
    )


    # =====================================================
    # RESOLUTION TREND
    # =====================================================

    story.append(
        Paragraph(
            "Resolution Trend",
            section_style
        )
    )


    resolution_table_data = [
        [
            "Month",
            "Filed",
            "Resolved",
        ]
    ]


    for (
        label,
        filed,
        resolved
    ) in zip(
        data["resolution_labels"],
        data["resolution_filed"],
        data["resolution_resolved"]
    ):

        resolution_table_data.append(
            [
                label,
                filed,
                resolved,
            ]
        )


    resolution_table = Table(
        resolution_table_data,
        colWidths=[
            45 * mm,
            45 * mm,
            45 * mm,
        ],
        repeatRows=1,
        hAlign="LEFT"
    )


    apply_pdf_table_style(
        resolution_table
    )


    story.append(
        resolution_table
    )


    # =====================================================
    # EVIDENCE
    # =====================================================

    story.append(
        Paragraph(
            "Evidence / Blockchain",
            section_style
        )
    )


    evidence_table_data = [

        [
            "Metric",
            "Value",
        ],

        [
            "Total Evidence",
            data["total_evidence"],
        ],

        [
            "Registered",
            data["registered_evidence"],
        ],

        [
            "Pending",
            data["pending_evidence"],
        ],

        [
            "Failed",
            data["failed_evidence"],
        ],

        [
            "Registration Rate",
            f"{data['blockchain_rate']}%",
        ],
    ]


    evidence_table = Table(
        evidence_table_data,
        colWidths=[
            90 * mm,
            45 * mm,
        ],
        repeatRows=1,
        hAlign="LEFT"
    )


    apply_pdf_table_style(
        evidence_table
    )


    story.append(
        evidence_table
    )


    # =====================================================
    # OFFICIAL ACTIVITY
    # =====================================================

    story.append(
        Paragraph(
            "Official Activity",
            section_style
        )
    )


    official_table_data = [
        [
            "Username",
            "Role",
            "Status",
            "Activities",
            "Last Activity",
        ]
    ]


    for officer in data[
        "official_activity"
    ]:

        last_activity = (
            officer["last_activity"]
            if officer["last_activity"]
            else "-"
        )


        official_table_data.append(
            [
                officer["username"],
                str(
                    officer["role"]
                ).title(),
                (
                    "Active"
                    if officer["is_active"]
                    else "Inactive"
                ),
                officer["activity_count"],
                str(last_activity),
            ]
        )


    official_table = Table(
        official_table_data,
        colWidths=[
            45 * mm,
            35 * mm,
            30 * mm,
            30 * mm,
            65 * mm,
        ],
        repeatRows=1,
        hAlign="LEFT"
    )


    apply_pdf_table_style(
        official_table
    )


    story.append(
        official_table
    )


    document.build(
        story
    )


    pdf = buffer.getvalue()

    buffer.close()


    filename = (
        "bantaybarangay_reports_"
        f"{data['generated_date']:%Y-%m-%d}"
        ".pdf"
    )


    response = HttpResponse(
        pdf,
        content_type="application/pdf"
    )


    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )


    return response


# =========================================================
# PDF TABLE STYLE
# =========================================================

def apply_pdf_table_style(table):

    table.setStyle(
        TableStyle(
            [

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#064E3B"
                    )
                ),

                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8.5
                ),

                (
                    "ALIGN",
                    (0, 0),
                    (-1, 0),
                    "CENTER"
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),

                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#D1D5DB"
                    )
                ),

                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#F7FAF8"
                        ),
                    ]
                ),

                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),

                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),

            ]
        )
    )